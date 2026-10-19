"""Turn uploaded files into plain UTF-8 text. The parser is picked by extension, then by MIME type."""
from __future__ import annotations

import re
import unicodedata
from html.parser import HTMLParser
from pathlib import Path

import docx
import pandas as pd
import pypdf

from .chunking import SHEET_PREFIX  # chunking repeats this line + the header row per table chunk

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".xlsx", ".json", ".html", ".htm"}

# Fallback when the file name has no usable extension (uploads from a browser often do).
_MIME_TO_EXT = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "text/csv": ".csv",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "text/html": ".html",
    "application/json": ".json",
}

# Tags that should end a line when converting HTML to text.
_BLOCK_TAGS = {"p", "div", "br", "li", "tr", "table", "section", "article", "header", "footer",
               "blockquote", "pre", "h1", "h2", "h3", "h4", "h5", "h6"}


class UnsupportedFileType(ValueError):
    """Extension and MIME type are both unknown."""


class ParseError(RuntimeError):
    """File is corrupt or could not be read."""


class _TextExtractor(HTMLParser):
    """Collects visible text; drops <script> and <style> content."""

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"script", "style"}:
            self._skip += 1
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self._skip:
            self._skip -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


# A line that starts a list item, numbered heading or legal paragraph keeps its line break:
# "- x", "• x", "– x", "§ 3", "(1)", "(a)", "1. x", "2.3 Titel", "a) x".
_LIST_START = re.compile(
    r"^(?:[-–—*]\s|[•▪●◦·]|§|\(\w{1,4}\)|\d{1,3}\.(?:\d{1,3}\.?)*\s|[a-zA-Z]\)\s)"
)
# Hyphen at a line end right after a letter ("Prü-"), not a free-standing dash ("Beginn -").
_HYPHEN_END = re.compile(r"[^\W\d_]-$")
# "Lehr-" + "und Lernformen" is a suspended compound, not a split word: keep "Lehr- und".
_SUSPENDED = re.compile(r"^(?:und|oder|bzw\.|sowie|bis)\b")
_LIGATURES = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl"}


def _join_lines(lines: list[str]) -> str:
    out = lines[0]
    for line in lines[1:]:
        if _LIST_START.match(line):
            out += "\n" + line
        elif _HYPHEN_END.search(out):
            if _SUSPENDED.match(line):
                out += " " + line  # "Lehr- und Lernformen"
            elif line[:1].islower():
                out = out[:-1] + line  # "Prü-" + "fungsordnung"
            else:
                out += line  # "Bachelor-" + "Studiengang" keeps the hyphen
        else:
            out += " " + line
    return out


def normalize_pdf_text(text: str) -> str:
    """Undo pypdf's hard line wrapping in (German) PDFs.

    Hyphenated line ends are rejoined (the hyphen is dropped only before a lowercase continuation),
    wrapped lines inside a paragraph are joined with a space, blank-line paragraph breaks are kept,
    and lines that start a list item, numbered heading or § stay on their own line.
    """
    # pdfium marks a line-end hyphen as U+FFFE (often without the newline); restore "-\n" for the rules below.
    text = re.sub("￾\r?\n?", "-\n", text)
    text = unicodedata.normalize("NFC", text)  # decomposed umlauts (a + U+0308) -> ä
    for lig, repl in _LIGATURES.items():
        text = text.replace(lig, repl)
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("­", "").replace(" ", " ")
    paragraphs = []
    for block in re.split(r"\n[ \t]*\n", text):
        lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in block.split("\n")]
        lines = [ln for ln in lines if ln]
        if lines:
            paragraphs.append(_join_lines(lines))
    return "\n\n".join(paragraphs)


def parse_pdf_pages(path: str) -> list[tuple[int, str]]:
    """Return (1-based page number, normalized text) for every page that has text."""
    try:
        raw = _pdfium_pages(path)
    except Exception:
        # pdfium is preferred: pypdf splits words in some German PDFs ("Ges amtumfang").
        try:
            reader = pypdf.PdfReader(path)
            raw = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:  # pypdf raises many unrelated error types on corrupt files
            raise ParseError(f"could not read {Path(path).name}: {exc}") from exc
    pages = []
    for n, text in enumerate(raw, start=1):
        text = normalize_pdf_text(text)
        if text:
            pages.append((n, text))
    return pages


def _pdfium_pages(path: str) -> list[str]:
    import pypdfium2 as pdfium  # optional dependency; parse_pdf_pages falls back to pypdf

    doc = pdfium.PdfDocument(path)
    try:
        return [doc[i].get_textpage().get_text_range().replace("\r\n", "\n") for i in range(len(doc))]
    finally:
        doc.close()


def _pdf(path: str) -> str:
    reader = pypdf.PdfReader(path)
    pages = (page.extract_text() or "" for page in reader.pages)
    # Skip blank pages so an empty PDF gives "" instead of a run of separators.
    return "\n\n".join(p for p in pages if p.strip())


def _docx(path: str) -> str:
    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _cell(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # pandas turns integer columns with gaps into floats
    return str(value).strip()




def _table(name: str, df: pd.DataFrame) -> str:
    """Compact table block: `# Sheet: name`, a `col | col` header row, then one `v | v` line per row."""
    rows = []
    for rec in df.itertuples(index=False):
        cells = [_cell(v).replace("|", "/").replace("\n", " ") for v in rec]
        if any(cells):
            rows.append(" | ".join(cells))
    if not rows:
        return ""
    header = " | ".join(str(c).strip().replace("|", "/") for c in df.columns)
    return f"{SHEET_PREFIX}{name}\n{header}\n" + "\n".join(rows)


def _csv(path: str) -> str:
    try:
        df = pd.read_csv(path, encoding="utf-8", encoding_errors="replace")
    except pd.errors.EmptyDataError:
        return ""
    return _table(Path(path).stem, df)


def _xlsx(path: str) -> str:
    sheets = pd.read_excel(path, sheet_name=None)  # dict of sheet name -> DataFrame
    blocks = [_table(str(name), df) for name, df in sheets.items()]
    return "\n\n".join(b for b in blocks if b)


def _html(path: str) -> str:
    extractor = _TextExtractor()
    extractor.feed(Path(path).read_text(encoding="utf-8", errors="replace"))
    text = "".join(extractor.parts)
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n\s*\n+", "\n\n", text)


def _text(path: str) -> str:
    return Path(path).read_text(encoding="utf-8-sig", errors="replace")


_PARSERS = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".csv": _csv,
    ".xlsx": _xlsx,
    ".html": _html,
    ".htm": _html,
    ".txt": _text,
    ".md": _text,
    ".json": _text,
}


def _extension(path: str, file_type: str | None) -> str:
    ext = Path(path).suffix.lower()
    if ext in SUPPORTED_EXTENSIONS:
        return ext
    mime = (file_type or "").split(";")[0].strip().lower()  # drop "; charset=..." parameters
    if mime in _MIME_TO_EXT:
        return _MIME_TO_EXT[mime]
    raise UnsupportedFileType(f"unsupported file type: {Path(path).name} ({mime or 'no MIME type'})")


def parse_document(path: str, file_type: str | None = None) -> str:
    """Return the document's text. Empty documents give "". Corrupt files raise ParseError."""
    parser = _PARSERS[_extension(path, file_type)]
    try:
        return parser(path).strip()
    except Exception as exc:  # pypdf, python-docx and pandas raise many unrelated error types
        raise ParseError(f"could not read {Path(path).name}: {exc}") from exc
