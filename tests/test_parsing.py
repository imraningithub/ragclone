import pandas as pd
import pypdf
import pytest

from app.parsing import (ParseError, UnsupportedFileType, normalize_pdf_text, parse_document,
                         parse_pdf_pages)


def _blank_pdf(path):
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as f:
        writer.write(f)
    return str(path)


def test_blank_pdf_returns_empty_string(tmp_path):
    assert parse_document(_blank_pdf(tmp_path / "blank.pdf")) == ""


def test_corrupt_pdf_raises_parse_error(tmp_path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"this is not a pdf")
    with pytest.raises(ParseError):
        parse_document(str(bad))


def test_docx_paragraphs_and_table(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello paragraph")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "name"
    table.cell(0, 1).text = "score"
    table.cell(1, 0).text = "ann"
    table.cell(1, 1).text = "42"
    path = tmp_path / "doc.docx"
    doc.save(path)

    text = parse_document(str(path))
    assert "Hello paragraph" in text
    assert "name | score" in text
    assert "ann | 42" in text


def test_csv_compact_table(tmp_path):
    path = tmp_path / "people.csv"
    pd.DataFrame({"name": ["Ann", "B|ob"], "age": [30, None]}).to_csv(path, index=False)

    lines = parse_document(str(path)).splitlines()
    assert lines == ["# Sheet: people", "name | age", "Ann | 30", "B/ob |"]


def test_xlsx_one_block_per_sheet(tmp_path):
    path = tmp_path / "book.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame({"item": ["pen"], "qty": [3]}).to_excel(writer, sheet_name="Stock", index=False)
        pd.DataFrame({"city": ["Oslo"]}).to_excel(writer, sheet_name="Trips", index=False)

    text = parse_document(str(path))
    assert "# Sheet: Stock\nitem | qty\npen | 3" in text
    assert "# Sheet: Trips\ncity\nOslo" in text


def test_html_strips_tags_scripts_and_styles(tmp_path):
    path = tmp_path / "page.html"
    path.write_text(
        "<html><head><style>p{color:red}</style><script>var x = 1;</script></head>"
        "<body><h1>Title</h1><p>First &amp; last</p></body></html>",
        encoding="utf-8",
    )
    text = parse_document(str(path))
    assert "Title" in text
    assert "First & last" in text
    assert "var x" not in text
    assert "color:red" not in text
    assert "<" not in text


def test_plain_text_markdown_and_json(tmp_path):
    for name, body in [("a.txt", "plain words"), ("b.md", "# heading"), ("c.json", '{"k": 1}')]:
        path = tmp_path / name
        path.write_text(body, encoding="utf-8")
        assert parse_document(str(path)) == body


def test_empty_text_file_returns_empty_string(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")
    assert parse_document(str(path)) == ""


def test_mime_type_used_when_extension_unknown(tmp_path):
    path = tmp_path / "upload.bin"
    path.write_text("from mime", encoding="utf-8")
    assert parse_document(str(path), "text/plain; charset=utf-8") == "from mime"


def test_unsupported_extension_raises(tmp_path):
    path = tmp_path / "tool.exe"
    path.write_bytes(b"MZ")
    with pytest.raises(UnsupportedFileType):
        parse_document(str(path))


def test_unknown_extension_and_mime_raises(tmp_path):
    path = tmp_path / "data.xyz"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(UnsupportedFileType):
        parse_document(str(path), "application/octet-stream")


# ------------------------------------------------------------- KB-mode PDFs
def test_normalize_hyphen_before_lowercase_is_removed():
    assert normalize_pdf_text("Die Prü-\nfungsordnung gilt.") == "Die Prüfungsordnung gilt."


def test_normalize_hyphen_before_uppercase_is_kept():
    assert normalize_pdf_text("im Bachelor-\nStudiengang") == "im Bachelor-Studiengang"


def test_normalize_suspended_compound_keeps_space():
    assert normalize_pdf_text("Lehr-\nund Lernformen") == "Lehr- und Lernformen"


def test_normalize_joins_wrapped_lines_and_whitespace():
    assert normalize_pdf_text("Das Semester\n  beginnt   im\nOktober.\r\n") == "Das Semester beginnt im Oktober."


def test_normalize_keeps_list_items_and_headings():
    text = ("Voraussetzungen sind:\n- Abitur\n• Sprachnachweis\n– Praktikum\n§ 3 Prüfungen\n"
            "(1) Die Prüfung\nwird schriftlich abgelegt.\n2. Wiederholung\na) einmal\n4.1 Fristen")
    assert normalize_pdf_text(text).split("\n") == [
        "Voraussetzungen sind:", "- Abitur", "• Sprachnachweis", "– Praktikum", "§ 3 Prüfungen",
        "(1) Die Prüfung wird schriftlich abgelegt.", "2. Wiederholung", "a) einmal", "4.1 Fristen",
    ]


def test_normalize_keeps_paragraphs_and_plain_numbers():
    text = "Erster Absatz\nzweite Zeile.\n\n \nZweiter Absatz im\n2025 Jahr."
    assert normalize_pdf_text(text) == "Erster Absatz zweite Zeile.\n\nZweiter Absatz im 2025 Jahr."


def test_normalize_umlauts_and_unicode_cleanup():
    decomposed = "Größe Prüfung"  # "ö"/"ü" as base letter + combining diaeresis
    assert normalize_pdf_text(decomposed) == "Größe Prüfung"
    assert normalize_pdf_text("Zu­lassung ﬁnden Äußerung") == "Zulassung finden Äußerung"


def _multipage_pdf(path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    pdf.add_page()
    pdf.multi_cell(0, 8, "Die Prü-\nfungsordnung regelt\ndas Studium.")
    pdf.add_page()  # blank page 2
    pdf.add_page()
    pdf.multi_cell(0, 8, "Größe und Maß\n- erster Punkt")
    pdf.output(str(path))
    return str(path)


def test_parse_pdf_pages_numbers_and_skips_blank(tmp_path):
    pages = parse_pdf_pages(_multipage_pdf(tmp_path / "ordnung.pdf"))
    assert [n for n, _ in pages] == [1, 3]
    assert pages[0][1] == "Die Prüfungsordnung regelt das Studium."
    assert "Größe und Maß\n- erster Punkt" == pages[1][1]


def test_parse_pdf_pages_blank_and_corrupt(tmp_path):
    assert parse_pdf_pages(_blank_pdf(tmp_path / "blank.pdf")) == []
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"not a pdf")
    with pytest.raises(ParseError):
        parse_pdf_pages(str(bad))


def test_parse_document_pdf_is_not_normalized(tmp_path):
    text = parse_document(_multipage_pdf(tmp_path / "raw.pdf"))
    assert "Prü-\nfungsordnung" in text
